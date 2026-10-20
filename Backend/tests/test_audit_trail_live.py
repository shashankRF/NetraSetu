from io import BytesIO
import openpyxl
import pytest
from fastapi.testclient import TestClient

from Backend.server import app
from Backend.audit.recorder import (
    create_audit_record,
    append_audit_record,
    load_audit_records,
)
from Backend.audit_api import build_audit_excel_workbook, _sanitize_for_excel


@pytest.fixture
def client():
    return TestClient(app)


def test_audit_list_endpoint(client):
    res = client.get("/api/audits")
    assert res.status_code == 200
    data = res.json()
    assert "total" in data
    assert "metrics" in data
    assert "audits" in data
    assert isinstance(data["audits"], list)
    if data["audits"]:
        first = data["audits"][0]
        assert "audit_id" in first
        assert "execution_id" in first
        assert "execution_status" in first
        assert "provenance" in first


def test_excel_export_endpoint(client):
    res = client.get("/api/audits/export/excel")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.headers["content-disposition"].endswith('.xlsx"')

    # Verify genuine Excel workbook structure
    wb = openpyxl.load_workbook(BytesIO(res.content))
    assert "HackGrid Audit Trail" in wb.sheetnames
    ws = wb["HackGrid Audit Trail"]
    headers = [cell.value for cell in ws[1]]
    assert "Audit ID" in headers
    assert "Execution / Run ID" in headers
    assert "Event Type" in headers
    assert "Execution Status" in headers
    assert "Error Details" in headers


def test_empty_log_excel_export_graceful():
    empty_bytes = build_audit_excel_workbook([])
    wb = openpyxl.load_workbook(BytesIO(empty_bytes))
    ws = wb.active
    assert ws.title == "HackGrid Audit Trail"
    assert ws.cell(row=1, column=1).value == "Audit ID"
    note_val = ws.cell(row=2, column=1).value
    assert "No audit events recorded yet" in str(note_val)


def test_formula_injection_sanitization():
    unsafe_strings = [
        "=cmd|' /C calc'!A0",
        "+1234567890",
        "-SUM(A1:A10)",
        "@HYPERLINK('http://malicious.com')",
        "\tmalicious_tab",
        "\rmalicious_cr",
        "\nmalicious_newline",
    ]
    for s in unsafe_strings:
        sanitized = _sanitize_for_excel(s)
        assert sanitized.startswith("'"), f"Formula injection not escaped: {repr(s)} -> {repr(sanitized)}"

    safe_string = "Normal security text"
    assert _sanitize_for_excel(safe_string) == safe_string


def test_query_execution_creates_audit_record(client):
    q_text = "What is the primary role of CERT-In under Indian cybersecurity regulations?"
    res = client.post("/api/query", json={"question": q_text, "jurisdiction": "India", "demo_mode": True})
    assert res.status_code == 200
    res_data = res.json()
    audit_id = res_data.get("audit_id")
    exec_id = res_data.get("execution_id")
    assert audit_id is not None
    assert exec_id is not None

    records = load_audit_records()
    matched = [r for r in records if r.get("audit_id") == audit_id]
    assert len(matched) == 1
    rec = matched[0]
    assert rec["question"] == q_text
    assert rec["execution_id"] == exec_id
    assert rec["execution_status"] == "COMPLETED"
    assert rec["event_type"] == "QUESTION_ANALYSIS"
    assert "Retrieval" in rec.get("workflow_stages", [])


def test_failed_query_creates_failure_audit_record(client):
    records_before = len(load_audit_records())
    res = client.post("/api/query", json={"question": "Valid query text?", "jurisdiction": "Atlantis"})
    assert res.status_code == 400

    records_after = load_audit_records()
    assert len(records_after) > records_before
    latest = records_after[-1]
    assert latest["execution_status"] == "FAILED"
    assert "error_details" in latest
    assert latest["final_decision"] == "VALIDATION_FAILED"


def test_scenario_execution_audit_record(client):
    res = client.post("/api/scenarios/run", json={"scenario_id": "HG-5G-003", "step_id": "T1", "jurisdiction": "India"})
    assert res.status_code == 200
    data = res.json()
    pipe_res = data.get("pipeline_result", {})
    audit_id = pipe_res.get("audit_id")
    assert audit_id is not None

    records = load_audit_records()
    matched = [r for r in records if r.get("audit_id") == audit_id]
    assert len(matched) == 1
    rec = matched[0]
    assert rec["event_type"] == "SCENARIO_EXECUTION"
    assert rec["scenario_id"] == "HG-5G-003"
    assert rec["step_id"] == "T1"


def test_governance_decision_creates_linked_audit_record(client):
    test_rec = create_audit_record(
        question="Should high-risk critical infrastructure slice be deployed?",
        human_review_required=True,
        final_decision="PENDING_HUMAN_REVIEW",
        risk_assessment={"risk_level": "CRITICAL", "risk_score": 90, "human_review_required": True},
    )
    audit_id = append_audit_record(test_rec)
    exec_id = test_rec["execution_id"]

    review_res = client.post(
        f"/api/audits/{audit_id}/review",
        json={"action": "APPROVED", "reviewer": "Senior Authority Officer"},
    )
    assert review_res.status_code == 200
    rev_data = review_res.json()
    assert rev_data["human_action"] == "APPROVED"
    gov_audit_id = rev_data.get("governance_audit_id")
    assert gov_audit_id is not None

    detail_res = client.get(f"/api/audits/{audit_id}")
    assert detail_res.status_code == 200
    detail_data = detail_res.json()
    assert detail_data["human_action"] == "APPROVED"
    assert len(detail_data.get("related_events", [])) >= 1
    related = detail_data["related_events"][0]
    assert related["execution_id"] == exec_id
    assert related["event_type"] == "GOVERNANCE_DECISION"
